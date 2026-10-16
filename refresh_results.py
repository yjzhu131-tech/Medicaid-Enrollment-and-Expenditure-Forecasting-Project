"""Rebuild clean data, every forecast/chart, and the advisor figure document."""
from pathlib import Path
import subprocess
import sys

from docx import Document

ROOT = Path(__file__).resolve().parent


def main():
    scripts = [ROOT / "clean_enrollment.py", ROOT / "clean_expenditure.py"]
    scripts.extend(sorted(ROOT.glob("*/code/*.py")))
    for script in scripts:
        subprocess.run([sys.executable, str(script)], cwd=ROOT, check=True)

    document_path = ROOT / "State_Result_Figures_for_Advisor.docx"
    document = Document(document_path)
    # Keep the existing document layout, captions, image sizes and ordering.
    for paragraph in document.paragraphs:
        if paragraph.text == "State Forecast Result Figures":
            paragraph.text = "State Forecast Result Figures — Medicaid + CHIP"
    images = iter(document.inline_shapes)
    for paragraph in document.paragraphs:
        if paragraph.text.strip().endswith(".png"):
            shape = next(images)
            relationship_id = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
            document.part.related_parts[relationship_id]._blob = (ROOT / paragraph.text.strip()).read_bytes()
    document.save(document_path)
    print("Updated all results and advisor document.")


if __name__ == "__main__":
    main()
